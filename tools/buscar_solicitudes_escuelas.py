#!/usr/bin/env python3
"""
Genera reportes de solicitudes por escuelas a partir de un respaldo SQL diario.

Flujo:
1. Importa el SQL en una base temporal de MariaDB/MySQL.
2. Carga escuelas objetivo desde CSV.
3. Cruza solicitudes por CCT, nombre, alias, localidad y municipio.
4. Anexa comentarios de seguimiento y tareas.
5. Genera CSV, Markdown y DOCX.
6. Elimina la base temporal.
"""

import argparse
import csv
import html
import json
import re
import subprocess
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path


DEFAULT_MYSQL = "/Applications/XAMPP/xamppfiles/bin/mysql"
STATUS_ORDER = ["No iniciada", "En proceso", "Concluida", "Pendiente"]


def clean(value):
    if value is None:
        return ""
    text = html.unescape(str(value))
    text = text.replace("\\r\\n", "\n").replace("\\n", "\n")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def norm(value):
    text = clean(value)
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return text.upper()


def regex_phrase(value):
    words = re.findall(r"[A-Z0-9]+", norm(value))
    if not words:
        return None
    return r"\b" + r"[^A-Z0-9]+".join(re.escape(word) for word in words) + r"\b"


def split_semicolon(value):
    return [part.strip() for part in clean(value).split(";") if part.strip()]


def split_entries(value):
    text = clean(value)
    if not text:
        return []
    return [part.strip() for part in text.split("\n---\n") if part.strip()]


def mysql_run(mysql, sql=None, db=None, input_file=None, capture=False):
    cmd = [mysql, "-uroot"]
    if db:
        cmd.append(db)
    if sql is not None:
        cmd.extend(["--default-character-set=utf8mb4", "-e", sql])
    stdin = None
    if input_file is not None:
        stdin = open(input_file, "rb")
    try:
        return subprocess.run(
            cmd,
            stdin=stdin,
            capture_output=capture,
            text=capture,
            check=True,
        )
    finally:
        if stdin is not None:
            stdin.close()


def mysql_json(mysql, db, inner_sql):
    sql = f"SELECT HEX(CONVERT(j USING utf8mb4)) FROM ({inner_sql}) q"
    result = subprocess.run(
        [
            mysql,
            "-uroot",
            db,
            "--default-character-set=utf8mb4",
            "--batch",
            "--raw",
            "--skip-column-names",
            "-e",
            sql,
        ],
        capture_output=True,
        text=True,
        check=True,
    )
    rows = []
    for line in result.stdout.splitlines():
        hex_text = line.strip()
        if hex_text:
            rows.append(json.loads(bytes.fromhex(hex_text).decode("utf-8")))
    return rows


def load_schools(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    schools = []
    for row in rows:
        cct = clean(row["cct"]).upper()
        names = [row.get("nombre", ""), *split_semicolon(row.get("alias", ""))]
        name_patterns = [regex_phrase(name) for name in names if clean(name)]
        loc_patterns = [
            regex_phrase(row.get("localidad", "")),
            regex_phrase(row.get("municipio", "")),
        ]
        schools.append(
            {
                "cct": cct,
                "nombre": clean(row.get("nombre", "")),
                "municipio": clean(row.get("municipio", "")),
                "localidad": clean(row.get("localidad", "")),
                "name_patterns": [p for p in name_patterns if p],
                "loc_patterns": [p for p in loc_patterns if p],
            }
        )
    return schools


def text_near(text, name_patterns, loc_patterns, window=140):
    for pattern in name_patterns:
        for match in re.finditer(pattern, text):
            chunk = text[max(0, match.start() - window) : match.end() + window]
            if any(re.search(loc_pattern, chunk) for loc_pattern in loc_patterns):
                return True
    return False


def import_sql(mysql, sql_path, db_name):
    mysql_run(mysql, f"DROP DATABASE IF EXISTS `{db_name}`; CREATE DATABASE `{db_name}` CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;")
    mysql_run(mysql, db=db_name, input_file=sql_path)


def fetch_data(mysql, db_name):
    schools = mysql_json(
        mysql,
        db_name,
        """
        SELECT JSON_OBJECT(
          'cct', CLAVECCT, 'nombre', NOMBRECT, 'nivel', N_NIVEL,
          'localidad', N_LOCALIDAD, 'municipio', N_MUNICIPIO,
          'domicilio', DOMICILIO, 'director', CONCAT_WS(' ', DIRECTOR, APELLIDO1, APELLIDO2)
        ) j
        FROM cct
        """,
    )
    solicitudes = mysql_json(
        mysql,
        db_name,
        """
        SELECT JSON_OBJECT(
          'id', s.id, 'folio', s.folio, 'categoria', s.fk_categoria,
          'cct', s.cct, 'cct_nombre', c.NOMBRECT, 'cct_localidad', c.N_LOCALIDAD,
          'cct_municipio', c.N_MUNICIPIO,
          'solicitante', CONCAT_WS(' ', s.solicitante, s.ap1, s.ap2),
          'fecha_peticion', s.fecha_peticion, 'tipo_procedencia', s.tipo_procedencia,
          'compromiso', s.compromiso, 'telefono', s.telefono, 'descripcion', s.descripcion,
          'indicaciones_secretario', s.indicaciones_secretario,
          'prioridad', s.prioridad, 'estado', s.estado, 'fecha_creacion', s.fecha_creacion
        ) j
        FROM solicitudes s
        LEFT JOIN cct c ON c.CLAVECCT = s.cct
        WHERE s.eliminado = 0
        ORDER BY s.fecha_peticion, s.id
        """,
    )
    comentarios = mysql_json(
        mysql,
        db_name,
        """
        SELECT JSON_OBJECT(
          'id', cs.id, 'solicitud_id', cs.solicitud_id,
          'fecha_comentario', cs.fecha_comentario, 'comentario', cs.comentario,
          'usuario_id', cs.usuario_id, 'usuario', u.nombre_completo,
          'archivo', cs.archivo, 'fecha_creacion', cs.fecha_creacion
        ) j
        FROM comentarios_seguimiento cs
        LEFT JOIN usuarios u ON u.id = cs.usuario_id
        ORDER BY cs.solicitud_id, cs.fecha_comentario, cs.id
        """,
    )
    tareas = mysql_json(
        mysql,
        db_name,
        """
        SELECT JSON_OBJECT(
          'pk_tarea', st.pk_tarea, 'fk_solicitud', st.fk_solicitud,
          'fk_usuario', st.fk_usuario, 'usuario', u.nombre_completo,
          'tarea', st.tarea, 'fecha', st.fecha, 'fecha_cierre', st.fecha_cierre,
          'estatus', CASE st.fk_estatus
            WHEN 0 THEN 'No iniciada'
            WHEN 1 THEN 'En proceso'
            WHEN 2 THEN 'Concluida'
            WHEN 3 THEN 'Pendiente'
            ELSE st.fk_estatus
          END
        ) j
        FROM solicitudes_tareas st
        LEFT JOIN usuarios u ON u.id = st.fk_usuario
        ORDER BY st.fk_solicitud, st.fecha, st.pk_tarea
        """,
    )
    return schools, solicitudes, comentarios, tareas


def build_matches(target_schools, solicitudes, comentarios, tareas):
    by_comment = defaultdict(list)
    by_task = defaultdict(list)
    for comment in comentarios:
        by_comment[comment["solicitud_id"]].append(comment)
    for task in tareas:
        by_task[task["fk_solicitud"]].append(task)

    matches = []
    for solicitud in solicitudes:
        text = " ".join(
            clean(solicitud.get(key, ""))
            for key in [
                "cct",
                "cct_nombre",
                "cct_localidad",
                "cct_municipio",
                "tipo_procedencia",
                "compromiso",
                "descripcion",
                "indicaciones_secretario",
            ]
        )
        text += " " + " ".join(clean(c.get("comentario", "")) for c in by_comment.get(solicitud["id"], []))
        text += " " + " ".join(clean(t.get("tarea", "")) for t in by_task.get(solicitud["id"], []))
        normalized = norm(text)
        solicitud_cct = norm(solicitud.get("cct", ""))

        for school in target_schools:
            reasons = []
            cct = school["cct"]
            if solicitud_cct == cct:
                reasons.append("CCT exacto en solicitud")
            if re.search(r"\b" + re.escape(cct) + r"\b", normalized):
                reasons.append("CCT mencionado en texto/seguimiento")
            if not reasons and text_near(normalized, school["name_patterns"], school["loc_patterns"]):
                reasons.append("nombre + localidad/municipio cercano")
            if not reasons:
                continue

            item = dict(solicitud)
            item.update(
                {
                    "escuela_cct": cct,
                    "escuela_nombre": school["nombre"],
                    "escuela_localidad": school["localidad"],
                    "escuela_municipio": school["municipio"],
                    "match_reason": "; ".join(dict.fromkeys(reasons)),
                    "comentarios": by_comment.get(solicitud["id"], []),
                    "tareas": by_task.get(solicitud["id"], []),
                }
            )
            matches.append(item)
    order = {school["cct"]: index for index, school in enumerate(target_schools)}
    matches.sort(key=lambda row: (order.get(row["escuela_cct"], 999), row.get("fecha_peticion") or "", row["id"]))
    return matches


def write_csv(matches, output_prefix):
    path = Path(f"{output_prefix}.csv")
    fields = [
        "escuela_cct",
        "escuela_nombre",
        "escuela_localidad",
        "escuela_municipio",
        "folio",
        "id",
        "cct",
        "cct_nombre",
        "cct_localidad",
        "cct_municipio",
        "fecha_peticion",
        "estado",
        "prioridad",
        "solicitante",
        "telefono",
        "tipo_procedencia",
        "compromiso",
        "descripcion",
        "indicaciones_secretario",
        "match_reason",
        "comentarios_seguimiento",
        "tareas_seguimiento",
    ]
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for match in matches:
            comments = "\n---\n".join(
                f"{clean(c.get('fecha_comentario'))} | {clean(c.get('usuario')) or 'usuario ' + str(c.get('usuario_id'))}: "
                f"{clean(c.get('comentario'))}"
                + (f" [archivo: {clean(c.get('archivo'))}]" if c.get("archivo") else "")
                for c in match["comentarios"]
            )
            task_text = "\n---\n".join(
                f"{clean(t.get('fecha'))} | {clean(t.get('usuario')) or 'usuario ' + str(t.get('fk_usuario'))} | "
                f"{clean(t.get('estatus'))}: {clean(t.get('tarea'))}"
                + (f" | cierre: {clean(t.get('fecha_cierre'))}" if t.get("fecha_cierre") else "")
                for t in match["tareas"]
            )
            row = {key: clean(match.get(key, "")) for key in fields if key not in {"comentarios_seguimiento", "tareas_seguimiento"}}
            row["comentarios_seguimiento"] = comments
            row["tareas_seguimiento"] = task_text
            writer.writerow(row)
    return path


def write_markdown(matches, target_schools, output_prefix, source_sql):
    path = Path(f"{output_prefix}.md")
    grouped = defaultdict(list)
    for match in matches:
        grouped[match["escuela_cct"]].append(match)

    lines = [
        "# Reporte de solicitudes por escuelas\n",
        f"Fuente: `{source_sql}`  \n",
        "Criterio: solicitudes no eliminadas con CCT exacto, CCT mencionado, o nombre de escuela con localidad/municipio cercano en el mismo texto. Se excluyen homonimos de otros CCT/municipios.\n",
        f"Total de solicitudes: **{len({m['id'] for m in matches})}**. Comentarios: **{sum(len(m['comentarios']) for m in matches)}**. Tareas: **{sum(len(m['tareas']) for m in matches)}**.\n",
        "## Resumen por escuela\n",
    ]
    for school in target_schools:
        items = grouped.get(school["cct"], [])
        status_counts = Counter(item["estado"] for item in items)
        lines.append(
            f"- **{school['cct']} - {school['nombre']}**: {len(items)} solicitudes; "
            + ", ".join(f"{status}: {status_counts[status]}" for status in STATUS_ORDER)
            + "."
        )
    lines.append("\n## Detalle\n")
    for school in target_schools:
        items = grouped.get(school["cct"], [])
        if not items:
            continue
        lines.append(f"### {school['cct']} - {school['nombre']}\n")
        for item in items:
            lines.append(f"#### {clean(item['folio'])} (ID {item['id']}) - {clean(item.get('fecha_peticion'))}\n")
            lines.append(f"- Estado/prioridad: {clean(item.get('estado'))} / {clean(item.get('prioridad'))}")
            lines.append(f"- Solicitante: {clean(item.get('solicitante'))}; telefono: {clean(item.get('telefono')) or 's/d'}")
            lines.append(f"- Compromiso: {clean(item.get('compromiso')) or 's/d'}")
            lines.append(f"- Coincidencia: {clean(item.get('match_reason'))}")
            if clean(item.get("descripcion")):
                lines.append("\nSolicitud/descripcion:\n\n" + clean(item.get("descripcion")) + "\n")
            if clean(item.get("indicaciones_secretario")):
                lines.append("Indicaciones del secretario:\n\n" + clean(item.get("indicaciones_secretario")) + "\n")
            lines.append("Comentarios de seguimiento:")
            if item["comentarios"]:
                for comment in item["comentarios"]:
                    lines.append(
                        f"- {clean(comment.get('fecha_comentario'))} | {clean(comment.get('usuario'))}: {clean(comment.get('comentario'))}"
                    )
            else:
                lines.append("- Sin comentarios registrados.")
            lines.append("Tareas/seguimientos asignados:")
            if item["tareas"]:
                for task in item["tareas"]:
                    lines.append(
                        f"- {clean(task.get('fecha'))} | {clean(task.get('usuario'))} | {clean(task.get('estatus'))}: {clean(task.get('tarea'))}"
                    )
            else:
                lines.append("- Sin tareas registradas.")
            lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def require_docx():
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise SystemExit(
            "No se encontro python-docx. Ejecuta este script con el Python empaquetado de Codex "
            "o usa --no-docx para generar solo CSV/Markdown."
        ) from exc
    return Document, WD_ORIENT, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT, WD_ALIGN_PARAGRAPH, OxmlElement, qn, Inches, Pt, RGBColor


def write_docx(matches, target_schools, output_prefix, source_sql):
    (
        Document,
        WD_ORIENT,
        WD_CELL_VERTICAL_ALIGNMENT,
        WD_TABLE_ALIGNMENT,
        WD_ALIGN_PARAGRAPH,
        OxmlElement,
        qn,
        Inches,
        Pt,
        RGBColor,
    ) = require_docx()

    def set_cell_shading(cell, fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), fill)

    def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
            node = tc_mar.find(qn(f"w:{name}"))
            if node is None:
                node = OxmlElement(f"w:{name}")
                tc_mar.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def set_table_width(table, width_dxa=9360, indent_dxa=120):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(width_dxa))
        tbl_w.set(qn("w:type"), "dxa")
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        if tbl_ind is None:
            tbl_ind = OxmlElement("w:tblInd")
            tbl_pr.append(tbl_ind)
        tbl_ind.set(qn("w:w"), str(indent_dxa))
        tbl_ind.set(qn("w:type"), "dxa")

    def set_cell_text(cell, text, bold=False, color=None, size=9):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(clean(text) or "s/d")
        run.bold = bold
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    def configure_styles(doc):
        styles = doc.styles
        normal = styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25
        for name, size, color, before, after in [
            ("Heading 1", 16, "2E74B5", 18, 10),
            ("Heading 2", 13, "2E74B5", 14, 7),
            ("Heading 3", 12, "1F4D78", 10, 5),
        ]:
            style = styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor.from_string(color)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
        styles["Body Text"].font.name = "Calibri"
        styles["Body Text"].font.size = Pt(10.5)
        styles["Body Text"].paragraph_format.space_after = Pt(6)
        styles["Body Text"].paragraph_format.line_spacing = 1.2
        muted = styles.add_style("Muted", 1)
        muted.font.name = "Calibri"
        muted.font.size = Pt(9)
        muted.font.color.rgb = RGBColor(85, 85, 85)
        label = styles.add_style("Case Label", 1)
        label.font.name = "Calibri"
        label.font.size = Pt(10)
        label.font.bold = True
        label.font.color.rgb = RGBColor(31, 77, 120)
        label.paragraph_format.space_before = Pt(6)
        label.paragraph_format.space_after = Pt(3)
        label.paragraph_format.keep_with_next = True
        bullet = styles["List Bullet"]
        bullet.font.name = "Calibri"
        bullet.font.size = Pt(9.5)
        bullet.paragraph_format.left_indent = Inches(0.375)
        bullet.paragraph_format.first_line_indent = Inches(-0.188)
        bullet.paragraph_format.space_after = Pt(4)
        bullet.paragraph_format.line_spacing = 1.2

    def add_long_text(doc, title, text):
        text = clean(text)
        if not text:
            return
        heading = doc.add_paragraph()
        heading.style = "Case Label"
        heading.add_run(title)
        for block in text.split("\n"):
            block = block.strip()
            if block:
                paragraph = doc.add_paragraph(block)
                paragraph.style = "Body Text"

    def add_entries(doc, title, entries):
        heading = doc.add_paragraph()
        heading.style = "Case Label"
        heading.add_run(title)
        if not entries:
            paragraph = doc.add_paragraph("Sin registros.")
            paragraph.style = "Muted"
            return
        for entry in entries:
            paragraph = doc.add_paragraph()
            paragraph.style = "List Bullet"
            paragraph.add_run(f" {entry}")

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Reporte de solicitudes por escuelas")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)
    subtitle = doc.add_paragraph()
    subtitle.style = "Muted"
    subtitle.add_run(f"Fuente: {Path(source_sql).name} | Generado: {date.today().isoformat()}")
    note = doc.add_paragraph()
    note.style = "Body Text"
    note.add_run(
        "Criterio: solicitudes no eliminadas con CCT exacto, CCT mencionado, o nombre de escuela "
        "con localidad/municipio cercano en el mismo texto. Se excluyeron homonimos de otros CCT/municipios."
    )

    grouped = defaultdict(list)
    for match in matches:
        grouped[match["escuela_cct"]].append(match)

    doc.add_heading("Resumen", level=1)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table)
    headers = ["CCT", "Escuela", "Total", "No iniciada", "En proceso", "Concluida", "Pendiente"]
    widths = [0.85, 2.05, 0.55, 0.85, 0.8, 0.75, 0.65]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        set_cell_shading(cell, "E8EEF5")
        set_cell_text(cell, header, bold=True, color="0B2545", size=9)
    for school in target_schools:
        items = grouped.get(school["cct"], [])
        counts = Counter(item["estado"] for item in items)
        row = table.add_row()
        values = [school["cct"], school["nombre"], str(len(items)), *(str(counts[status]) for status in STATUS_ORDER)]
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            set_cell_text(cell, value, size=9)

    doc.add_heading("Detalle por escuela", level=1)
    for school in target_schools:
        items = grouped.get(school["cct"], [])
        if not items:
            continue
        doc.add_heading(f"{school['cct']} - {school['nombre']}", level=2)
        paragraph = doc.add_paragraph()
        paragraph.style = "Case Label"
        paragraph.add_run(f"Solicitudes localizadas: {len(items)}")
        for item in items:
            heading = doc.add_heading(f"{item['folio']} (ID {item['id']})", level=3)
            heading.paragraph_format.keep_with_next = True
            meta = doc.add_table(rows=4, cols=4)
            meta.style = "Table Grid"
            set_table_width(meta)
            pairs = [
                ("Fecha", item["fecha_peticion"]),
                ("Estado", item["estado"]),
                ("Prioridad", item["prioridad"]),
                ("CCT solicitud", item.get("cct") or "s/d"),
                ("Solicitante", item["solicitante"]),
                ("Telefono", item.get("telefono") or "s/d"),
                ("Procedencia", item.get("tipo_procedencia") or "s/d"),
                ("Coincidencia", item["match_reason"]),
            ]
            for index, (label, value) in enumerate(pairs):
                r_idx, c_idx = divmod(index, 2)
                label_cell = meta.rows[r_idx].cells[c_idx * 2]
                value_cell = meta.rows[r_idx].cells[c_idx * 2 + 1]
                set_cell_shading(label_cell, "F2F4F7")
                set_cell_text(label_cell, label, bold=True, color="1F4D78", size=8.5)
                set_cell_text(value_cell, value, size=8.5)
            add_long_text(doc, "Compromiso", item.get("compromiso"))
            add_long_text(doc, "Solicitud / descripcion", item.get("descripcion"))
            add_long_text(doc, "Indicaciones del secretario", item.get("indicaciones_secretario"))
            comments = [
                f"{clean(c.get('fecha_comentario'))} | {clean(c.get('usuario')) or 'usuario ' + str(c.get('usuario_id'))}: {clean(c.get('comentario'))}"
                + (f" [archivo: {clean(c.get('archivo'))}]" if c.get("archivo") else "")
                for c in item["comentarios"]
            ]
            tasks = [
                f"{clean(t.get('fecha'))} | {clean(t.get('usuario')) or 'usuario ' + str(t.get('fk_usuario'))} | {clean(t.get('estatus'))}: {clean(t.get('tarea'))}"
                + (f" | cierre: {clean(t.get('fecha_cierre'))}" if t.get("fecha_cierre") else "")
                for t in item["tareas"]
            ]
            add_entries(doc, "Comentarios de seguimiento", comments)
            add_entries(doc, "Tareas / seguimientos asignados", tasks)

    path = Path(f"{output_prefix}.docx")
    doc.core_properties.title = "Reporte de solicitudes por escuelas"
    doc.core_properties.subject = "Solicitudes y comentarios de seguimiento por escuelas seleccionadas"
    doc.core_properties.author = "Codex"
    doc.save(path)
    return path


def main():
    parser = argparse.ArgumentParser(description="Busca solicitudes por escuelas en un respaldo SQL diario.")
    parser.add_argument("--sql", required=True, help="Ruta del respaldo .sql descargado.")
    parser.add_argument("--escuelas", default="tools/escuelas_busqueda.csv", help="CSV con CCT/nombre/municipio/localidad/alias.")
    parser.add_argument("--salida", default="reporte_solicitudes_escuelas", help="Prefijo de salida sin extension.")
    parser.add_argument("--mysql", default=DEFAULT_MYSQL, help="Ruta al cliente mysql.")
    parser.add_argument("--db-temp", default="sesol_report_tmp", help="Nombre de base temporal.")
    parser.add_argument("--no-docx", action="store_true", help="Genera solo CSV y Markdown.")
    parser.add_argument("--conservar-db", action="store_true", help="No elimina la base temporal al terminar.")
    args = parser.parse_args()

    sql_path = Path(args.sql).expanduser().resolve()
    schools_path = Path(args.escuelas).expanduser()
    output_prefix = str(Path(args.salida).expanduser())
    target_schools = load_schools(schools_path)

    try:
        print("1/6 Importando SQL en base temporal...")
        import_sql(args.mysql, sql_path, args.db_temp)
        print("2/6 Leyendo tablas principales...")
        _all_schools, solicitudes, comentarios, tareas = fetch_data(args.mysql, args.db_temp)
        print("3/6 Cruzando por CCT, nombre, localidad y municipio...")
        matches = build_matches(target_schools, solicitudes, comentarios, tareas)
        print("4/6 Generando CSV y Markdown...")
        csv_path = write_csv(matches, output_prefix)
        md_path = write_markdown(matches, target_schools, output_prefix, sql_path)
        docx_path = None
        if not args.no_docx:
            print("5/6 Generando Word...")
            docx_path = write_docx(matches, target_schools, output_prefix, sql_path)
        else:
            print("5/6 Word omitido por --no-docx.")
        print("6/6 Limpieza...")
    finally:
        if not args.conservar_db:
            try:
                mysql_run(args.mysql, f"DROP DATABASE IF EXISTS `{args.db_temp}`;")
            except Exception as exc:
                print(f"Advertencia: no se pudo eliminar la base temporal {args.db_temp}: {exc}", file=sys.stderr)

    print(f"Solicitudes encontradas: {len({m['id'] for m in matches})}")
    print(f"CSV: {csv_path}")
    print(f"Markdown: {md_path}")
    if docx_path:
        print(f"Word: {docx_path}")


if __name__ == "__main__":
    main()
