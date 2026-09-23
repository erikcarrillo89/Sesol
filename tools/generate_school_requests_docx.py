import csv
import html
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Applications/XAMPP/xamppfiles/htdocs/sesol")
CSV_PATH = ROOT / "reporte_solicitudes_escuelas.csv"
OUT_PATH = ROOT / "reporte_solicitudes_escuelas.docx"


SCHOOL_ORDER = [
    "31DST0008U",
    "31EES0024S",
    "31ETV0123Y",
    "31EPR0133F",
    "31DJN0199Z",
    "31DPR0976X",
]


def clean(value):
    if value is None:
        return ""
    text = html.unescape(str(value))
    text = text.replace("\\r\\n", "\n").replace("\\n", "\n").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_entries(value):
    text = clean(value)
    if not text:
        return []
    return [part.strip() for part in text.split("\n---\n") if part.strip()]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=9360, indent_dxa=120):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(clean(text) or "s/d")
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_label_value(doc, label, value):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.color.rgb = RGBColor(31, 77, 120)
    paragraph.add_run(clean(value) or "s/d")


def add_long_text(doc, title, text):
    text = clean(text)
    if not text:
        return
    heading = doc.add_paragraph()
    heading.style = "Case Label"
    heading.add_run(title)
    for block in text.split("\n"):
        block = block.strip()
        if not block:
            continue
        para = doc.add_paragraph(block)
        para.style = "Body Text"


def add_entries(doc, title, entries):
    heading = doc.add_paragraph()
    heading.style = "Case Label"
    heading.add_run(title)
    if not entries:
        para = doc.add_paragraph("Sin registros.")
        para.style = "Muted"
        return
    for entry in entries:
        para = doc.add_paragraph()
        para.style = "List Bullet"
        para.paragraph_format.space_after = Pt(4)
        para.add_run(f" {entry}")


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["Body Text"].font.name = "Calibri"
    styles["Body Text"].font.size = Pt(10.5)
    styles["Body Text"].paragraph_format.space_after = Pt(6)
    styles["Body Text"].paragraph_format.line_spacing = 1.2

    muted = styles.add_style("Muted", 1)
    muted.font.name = "Calibri"
    muted.font.size = Pt(9)
    muted.font.color.rgb = RGBColor(85, 85, 85)
    muted.paragraph_format.space_after = Pt(6)

    label = styles.add_style("Case Label", 1)
    label.font.name = "Calibri"
    label.font.size = Pt(10)
    label.font.bold = True
    label.font.color.rgb = RGBColor(31, 77, 120)
    label.paragraph_format.space_before = Pt(6)
    label.paragraph_format.space_after = Pt(3)
    label.paragraph_format.keep_with_next = True

    bullet = styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(9.5)
    bullet.paragraph_format.left_indent = Inches(0.375)
    bullet.paragraph_format.first_line_indent = Inches(-0.188)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.2


def add_summary_table(doc, rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["escuela_cct"]].append(row)

    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table)
    headers = ["CCT", "Escuela", "Total", "No iniciada", "En proceso", "Concluida", "Pendiente"]
    widths = [0.85, 2.05, 0.55, 0.85, 0.8, 0.75, 0.65]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        set_cell_shading(cell, "E8EEF5")
        set_cell_text(cell, header, bold=True, color="0B2545", size=9)
    set_repeat_table_header(table.rows[0])

    for code in SCHOOL_ORDER:
        items = grouped.get(code, [])
        if not items:
            continue
        row = table.add_row()
        first = items[0]
        status_counts = defaultdict(int)
        for item in items:
            status_counts[item["estado"]] += 1
        values = [
            code,
            first["escuela_nombre"],
            str(len(items)),
            str(status_counts["No iniciada"]),
            str(status_counts["En proceso"]),
            str(status_counts["Concluida"]),
            str(status_counts["Pendiente"]),
        ]
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.width = Inches(widths[idx])
            set_cell_text(cell, value, size=9)


def add_metadata_table(doc, row):
    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table)
    pairs = [
        ("Fecha", row["fecha_peticion"]),
        ("Estado", row["estado"]),
        ("Prioridad", row["prioridad"]),
        ("CCT solicitud", row["cct"] or "s/d"),
        ("Solicitante", row["solicitante"]),
        ("Telefono", row["telefono"] or "s/d"),
        ("Procedencia", row["tipo_procedencia"]),
        ("Coincidencia", row["match_reason"]),
    ]
    for idx, (label, value) in enumerate(pairs):
        r_idx, c_idx = divmod(idx, 2)
        label_cell = table.rows[r_idx].cells[c_idx * 2]
        value_cell = table.rows[r_idx].cells[c_idx * 2 + 1]
        set_cell_shading(label_cell, "F2F4F7")
        set_cell_text(label_cell, label, bold=True, color="1F4D78", size=8.5)
        set_cell_text(value_cell, value, size=8.5)
    for row_cells in table.rows:
        row_cells.cells[0].width = Inches(1.0)
        row_cells.cells[1].width = Inches(2.25)
        row_cells.cells[2].width = Inches(1.0)
        row_cells.cells[3].width = Inches(2.25)


def main():
    with CSV_PATH.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Reporte de solicitudes por escuelas")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.style = "Muted"
    subtitle.add_run(f"Fuente: reporte_solicitudes_escuelas.csv | Generado: {date.today().isoformat()}")

    note = doc.add_paragraph()
    note.style = "Body Text"
    note.add_run(
        "Criterio: solicitudes no eliminadas con CCT exacto, CCT mencionado, o nombre de escuela "
        "con localidad/municipio cercano en el mismo texto. Se excluyeron homonimos de otros CCT/municipios."
    )

    doc.add_heading("Resumen", level=1)
    add_summary_table(doc, rows)

    doc.add_heading("Detalle por escuela", level=1)
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["escuela_cct"]].append(row)

    for code in SCHOOL_ORDER:
        items = grouped.get(code, [])
        if not items:
            continue
        doc.add_heading(f"{code} - {items[0]['escuela_nombre']}", level=2)
        add_label_value(doc, "Solicitudes localizadas", str(len(items)))
        for item in items:
            heading = doc.add_heading(f"{item['folio']} (ID {item['id']})", level=3)
            heading.paragraph_format.keep_with_next = True
            add_metadata_table(doc, item)
            add_long_text(doc, "Compromiso", item["compromiso"])
            add_long_text(doc, "Solicitud / descripcion", item["descripcion"])
            add_long_text(doc, "Indicaciones del secretario", item["indicaciones_secretario"])
            add_entries(doc, "Comentarios de seguimiento", split_entries(item["comentarios_seguimiento"]))
            add_entries(doc, "Tareas / seguimientos asignados", split_entries(item["tareas_seguimiento"]))

    doc.core_properties.title = "Reporte de solicitudes por escuelas"
    doc.core_properties.subject = "Solicitudes y comentarios de seguimiento por escuelas seleccionadas"
    doc.core_properties.author = "Codex"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
